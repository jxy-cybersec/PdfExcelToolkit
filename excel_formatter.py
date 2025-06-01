from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import os

def format_excel_with_highlights(file_path, sheet_name, column_letter, threshold):
    if not file_path.lower().endswith('.xlsx'):
        file_path += '.xlsx'
    if not os.path.exists(file_path):
        print("❌ Excel file not found.")
        return
    try:
        wb = load_workbook(file_path)
        if sheet_name not in wb.sheetnames:
            print("❌ Sheet name not found.")
            return

        ws = wb[sheet_name]
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        for row in ws.iter_rows(min_row=2):
            cell = row[ord(column_letter) - ord('A')]
            try:
                if cell.value is not None and float(cell.value) < threshold:
                    cell.fill = red_fill
            except:
                continue

        wb.save(file_path)
        print(f"✅ Formatting applied and saved to: {file_path}")
    except Exception as e:
        print(f"❌ Error: {e}")
